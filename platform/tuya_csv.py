"""Parse and validate Home Box Tuya Local device CSV / Excel rows."""

from __future__ import annotations

import base64
import csv
import io
import re
from typing import Any


REQUIRED = ("name", "device_id", "local_key", "host")
OPTIONAL_DEFAULTS = {
    "protocol_version": "auto",
    "type": "",
    "poll_only": "false",
}
CSV_FIELDS = list(REQUIRED) + list(OPTIONAL_DEFAULTS.keys())

HEADER_ALIASES = {
    "name": "name",
    "title": "name",
    "device_id": "device_id",
    "deviceid": "device_id",
    "id": "device_id",
    "local_key": "local_key",
    "localkey": "local_key",
    "key": "local_key",
    "host": "host",
    "ip": "host",
    "address": "host",
    "protocol_version": "protocol_version",
    "protocol": "protocol_version",
    "version": "protocol_version",
    "type": "type",
    "product": "type",
    "profile": "type",
    "poll_only": "poll_only",
    "poll": "poll_only",
}


def _norm_header(h: str) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_")
    return HEADER_ALIASES.get(key, key)


def _as_bool(value: str) -> bool:
    return str(value or "").strip().lower() in ("1", "true", "yes", "y", "on")


def parse_csv_text(text: str) -> tuple[list[dict[str, Any]], list[str]]:
    """Return (devices, errors)."""
    errors: list[str] = []
    raw = (text or "").lstrip("\ufeff").strip()
    if not raw:
        return [], ["CSV is empty"]

    reader = csv.DictReader(io.StringIO(raw))
    if not reader.fieldnames:
        return [], ["CSV has no header row"]

    mapping = {_norm_header(h): h for h in reader.fieldnames if h is not None}
    missing = [c for c in REQUIRED if c not in mapping]
    if missing:
        return [], [f"Missing required column(s): {', '.join(missing)}"]

    devices: list[dict[str, Any]] = []
    for i, row in enumerate(reader, start=2):
        item: dict[str, Any] = {}
        for canon, original in mapping.items():
            if canon in REQUIRED or canon in OPTIONAL_DEFAULTS:
                item[canon] = (row.get(original) or "").strip()
        for key, default in OPTIONAL_DEFAULTS.items():
            if not item.get(key):
                item[key] = default

        row_errors = []
        for key in REQUIRED:
            if not item.get(key):
                row_errors.append(key)
        if row_errors:
            errors.append(f"Row {i}: missing {', '.join(row_errors)}")
            continue

        item["poll_only"] = _as_bool(str(item.get("poll_only")))
        proto = str(item.get("protocol_version") or "auto").strip()
        item["protocol_version"] = proto
        devices.append(item)

    if not devices and not errors:
        errors.append("No data rows found")
    return devices, errors


def devices_to_csv(devices: list[dict[str, Any]]) -> str:
    """Serialize devices back to CSV text (for Excel → textarea round-trip)."""
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=CSV_FIELDS, extrasaction="ignore")
    writer.writeheader()
    for d in devices:
        row = {key: d.get(key, "") for key in CSV_FIELDS}
        row["poll_only"] = "true" if d.get("poll_only") else "false"
        writer.writerow(row)
    return buf.getvalue()


def parse_xlsx_bytes(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    """Parse first sheet of an .xlsx workbook into device rows."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [], [
            "Excel support not installed — rebuild home-box-tuya-import image"
        ]

    if not data:
        return [], ["Excel file is empty"]

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as err:  # noqa: BLE001
        return [], [f"Cannot read Excel file: {err}"]

    try:
        ws = wb.active
        if ws is None:
            return [], ["Excel workbook has no active sheet"]

        out = io.StringIO()
        writer = csv.writer(out)
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            writer.writerow(["" if c is None else str(c).strip() for c in row])
            row_count += 1
        if row_count == 0:
            return [], ["Excel sheet is empty"]
        return parse_csv_text(out.getvalue())
    finally:
        wb.close()


def parse_xlsx_base64(b64: str) -> tuple[list[dict[str, Any]], list[str]]:
    raw = (b64 or "").strip()
    if not raw:
        return [], ["Excel payload is empty"]
    try:
        data = base64.b64decode(raw, validate=False)
    except Exception as err:  # noqa: BLE001
        return [], [f"Invalid Excel base64: {err}"]
    return parse_xlsx_bytes(data)


def public_device(d: dict[str, Any]) -> dict[str, Any]:
    """Mask local_key for UI lists."""
    key = str(d.get("local_key") or "")
    masked = (key[:2] + "…" + key[-2:]) if len(key) > 6 else "…"
    return {
        "name": d.get("name"),
        "device_id": d.get("device_id"),
        "host": d.get("host"),
        "protocol_version": d.get("protocol_version"),
        "type": d.get("type") or "",
        "poll_only": bool(d.get("poll_only")),
        "local_key_masked": masked,
        "has_type": bool(d.get("type")),
    }
